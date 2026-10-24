#!/usr/bin/env python3
"""Übernimmt den Excel-Export (Blätter „Fahrzeuge“ und „Arbeitskleidung“) in personal/daten-lokal.js.
Die Datei bleibt nur auf diesem Rechner (.gitignore) – echte Daten nie ins Repo!
  python3 werkzeuge/excel_import.py Bruell_Mitarbeiter_Export.xlsx
Beim nächsten Öffnen des Dashboards werden die neuen Stände einmalig in den Browser übernommen.
"""
import json, re, sys, time, pathlib
import openpyxl

HIER = pathlib.Path(__file__).resolve().parent.parent
ZIEL = HIER / "personal" / "daten-lokal.js"
wb = openpyxl.load_workbook(sys.argv[1] if len(sys.argv) > 1 else HIER / "Bruell_Mitarbeiter_Export.xlsx")
txt = lambda v: "" if v is None or str(v).strip() in ("–", "-") else (v.strftime("%d.%m.%Y") if hasattr(v, "strftime") else str(v).strip())

src = ZIEL.read_text(encoding="utf8")
m = re.search(r"window\.PK_VOLL = (\{.*?\});\s*$", src, re.M | re.S)
pk = json.loads(m.group(1))
namen = {(d.get("vorname", "") + " " + d.get("nachname", "")).strip().lower(): nr for nr, d in pk.items()}

# --- Fahrzeuge ---
fz = {}
for r in wb["Fahrzeuge"].iter_rows(min_row=2, values_only=True):
    if r[0] is None: continue
    nr = str(int(r[0]))
    fz[nr] = {k: v for k, v in {"fz_nr_intern": nr.zfill(3), "fz_kennzeichen": txt(r[1]), "fz_modell": txt(r[2]), "fz_fahrer": txt(r[3]),
              "fz_typ": txt(r[4]), "fz_status": txt(r[5]), "fz_tuev": txt(r[6]), "fz_leasing_ende": txt(r[7])}.items() if v}

# --- Arbeitskleidung → _ak_table der Personalkarte ---
kleidung, fremd = {}, 0
for r in wb["Arbeitskleidung"].iter_rows(min_row=2, values_only=True):
    nr = namen.get(txt(r[0]).lower())
    if not nr: fremd += 1; continue
    kleidung.setdefault(nr, []).append([txt(v) for v in r[1:6]])
for nr, zeilen in kleidung.items(): pk[nr]["_ak_table"] = json.dumps(zeilen, ensure_ascii=False)

stand = time.strftime("%Y%m%d%H%M%S")
kopf = "/* VERTRAULICH – vollständige Personaldaten. Nur lokal, niemals ins Repo / auf GitHub (steht in .gitignore). */\n"
ZIEL.write_text(kopf + "window.PK_VOLL = " + json.dumps(pk, ensure_ascii=False) + ";\n"
                + "window.FZ_VOLL = " + json.dumps(fz, ensure_ascii=False) + ";\n"
                + f'window.DATEN_STAND = "{stand}";\n', encoding="utf8")
ZIEL.chmod(0o600)
print(f"✓ {len(fz)} Fahrzeuge, Kleidung für {len(kleidung)} Mitarbeiter ({sum(map(len, kleidung.values()))} Positionen)"
      + (f", {fremd} Zeilen ohne passende Personalkarte übersprungen" if fremd else "") + f" → {ZIEL.name} (Stand {stand})")
